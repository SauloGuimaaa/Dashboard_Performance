from __future__ import annotations

import logging
from dataclasses import asdict, is_dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Optional

import pandas as pd
from docx import Document
from docx.shared import Inches

logger = logging.getLogger(__name__)


def _safe_str(x: Any) -> str:
    if x is None:
        return ""
    return str(x)


def _truncate(text: str, max_len: int = 180) -> str:
    t = text.strip()
    if len(t) <= max_len:
        return t
    return t[: max_len - 1].rstrip() + "…"


def _add_kv_table(doc: Document, data: Dict[str, Any], title: str) -> None:
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = "Métrica"
    hdr[1].text = "Valor"

    for k, v in data.items():
        row = table.add_row().cells
        row[0].text = _safe_str(k)
        row[1].text = _safe_str(v)


def _add_df_table(doc: Document, df: pd.DataFrame, title: str, max_rows: int = 20) -> None:
    doc.add_heading(title, level=2)

    if df is None or len(df) == 0:
        doc.add_paragraph("Sem dados para esta seção.")
        return

    d = df.copy().head(max_rows)

    # evita colunas gigantes no DOCX
    for col in d.columns:
        if d[col].dtype == "object":
            d[col] = d[col].astype(str).map(lambda s: _truncate(s, 120))

    table = doc.add_table(rows=1, cols=len(d.columns))
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for j, col in enumerate(d.columns):
        hdr_cells[j].text = _safe_str(col)

    for _, row in d.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(d.columns):
            cells[j].text = _safe_str(row[col])


def build_report_docx(
    output_dir: Path,
    prefix: str,
    dataset_name: str,
    period_a_start: Optional[date],
    period_a_end: Optional[date],
    kpis_a: Any,
    top_reach: pd.DataFrame,
    top_eng: pd.DataFrame,
    daily: pd.DataFrame,
    weekly: pd.DataFrame,
    chart_paths: Dict[str, Path],
    notes: str = "",
) -> Path:
    """
    Gera um relatório DOCX em output/reports.

    - Inclui: KPIs, Top posts, tabelas de séries, e imagens (PNGs) se existirem.
    - chart_paths: dict com paths PNG (ex.: retornado do export_charts_png).
    """
    reports_dir = output_dir / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)

    safe_prefix = prefix.strip().replace(" ", "_")
    out_path = reports_dir / f"{safe_prefix}.docx"

    doc = Document()
    doc.add_heading("Relatório de Performance", level=1)
    doc.add_paragraph(f"Dataset: {dataset_name}")
    doc.add_paragraph(f"Gerado em: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    if period_a_start and period_a_end:
        doc.add_paragraph(f"Período (A): {period_a_start} → {period_a_end}")

    if notes.strip():
        doc.add_paragraph(f"Observações: {notes.strip()}")

    # KPIs
    if is_dataclass(kpis_a):
        kpis_dict = asdict(kpis_a)
    else:
        # fallback: tenta ler atributos comuns
        kpis_dict = {k: getattr(kpis_a, k) for k in dir(kpis_a) if not k.startswith("_")}

    _add_kv_table(doc, kpis_dict, title="KPIs (Período A)")

    # Gráficos (PNGs)
    doc.add_heading("Gráficos", level=1)
    if not chart_paths:
        doc.add_paragraph("Nenhum gráfico PNG disponível.")
    else:
        # ordenação estável (ajuda no relatório)
        for key in sorted(chart_paths.keys()):
            p = chart_paths[key]
            doc.add_heading(key, level=2)
            if p.exists():
                # 6.5" cabe bem em A4/Letter com margens padrão
                doc.add_picture(str(p), width=Inches(6.5))
            else:
                doc.add_paragraph(f"Arquivo não encontrado: {p}")

    # Top posts
    doc.add_page_break()
    doc.add_heading("Top Posts", level=1)
    _add_df_table(doc, top_reach, title="Top Reach", max_rows=20)
    _add_df_table(doc, top_eng, title="Top Engagement", max_rows=20)

    # Séries temporais
    doc.add_page_break()
    doc.add_heading("Evolução no tempo", level=1)
    _add_df_table(doc, daily, title="Série diária", max_rows=60)
    _add_df_table(doc, weekly, title="Série semanal", max_rows=60)

    doc.save(out_path)
    logger.info("DOCX report OK | %s | bytes=%d", str(out_path), out_path.stat().st_size)
    return out_path
